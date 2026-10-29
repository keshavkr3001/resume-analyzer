from io import BytesIO

from collections.abc import Generator

from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker

from app.database import Base, get_db, settings
from app.main import app
from app.models import Resume


def test_resume_text_is_only_visible_to_its_owner(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "upload_dir", str(tmp_path / "uploads"))
    engine = create_engine(f"sqlite:///{tmp_path / 'test.db'}", connect_args={"check_same_thread": False})
    Base.metadata.create_all(engine)
    test_session = sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)

    def override_get_db() -> Generator[Session, None, None]:
        with test_session() as session:
            yield session

    app.dependency_overrides[get_db] = override_get_db
    try:
        with TestClient(app) as client:
            owner_response = client.post(
                "/api/auth/register",
                json={"name": "Owner", "email": "owner@example.com", "password": "a-secure-password"},
            )
            other_response = client.post(
                "/api/auth/register",
                json={"name": "Other", "email": "other@example.com", "password": "another-secure-password"},
            )
            assert owner_response.status_code == 201
            assert other_response.status_code == 201
            owner = owner_response.json()
            other = other_response.json()
            with test_session() as session:
                resume = Resume(
                    user_id=owner["user"]["id"],
                    original_filename="resume.pdf",
                    stored_file_reference="test-file",
                    file_type="pdf",
                    extracted_text="A sufficiently long resume text for access control.",
                )
                session.add(resume)
                session.commit()
                resume_id = resume.id

            owner_result = client.get(
                f"/api/resumes/{resume_id}/text",
                headers={"Authorization": f"Bearer {owner['access_token']}"},
            )
            other_result = client.get(
                f"/api/resumes/{resume_id}/text",
                headers={"Authorization": f"Bearer {other['access_token']}"},
            )
            assert owner_result.status_code == 200
            assert owner_result.json()["text"].startswith("A sufficiently long")
            assert other_result.status_code == 404
            document = Document()
            document.add_paragraph("Education and skills")
            document.add_paragraph("A student resume containing Python, SQL, projects, and relevant experience.")
            buffer = BytesIO()
            document.save(buffer)
            upload_response = client.post(
                "/api/resumes/upload",
                headers={"Authorization": f"Bearer {owner['access_token']}"},
                files={"file": ("student.docx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )
            assert upload_response.status_code == 201
            uploaded_id = upload_response.json()["id"]
            uploaded_text = client.get(
                f"/api/resumes/{uploaded_id}/text",
                headers={"Authorization": f"Bearer {owner['access_token']}"},
            )
            assert "A student resume" in uploaded_text.json()["text"]
            invalid_upload = client.post(
                "/api/resumes/upload",
                headers={"Authorization": f"Bearer {owner['access_token']}"},
                files={"file": ("resume.txt", b"unsupported", "text/plain")},
            )
            assert invalid_upload.status_code == 415
            analysis_response = client.post(
                f"/api/resumes/{resume_id}/analyze",
                headers={"Authorization": f"Bearer {owner['access_token']}"},
            )
            assert analysis_response.status_code == 200
            assert analysis_response.json()["analysis_mode"] == "rule-based"
            report_response = client.get(
                f"/api/resumes/{resume_id}/report/json",
                headers={"Authorization": f"Bearer {owner['access_token']}"},
            )
            assert report_response.status_code == 200
            assert 0 <= report_response.json()["quality_score"] <= 100
            inaccessible_analysis = client.get(
                f"/api/resumes/{resume_id}/analysis",
                headers={"Authorization": f"Bearer {other['access_token']}"},
            )
            assert inaccessible_analysis.status_code == 404
            deletion_response = client.delete(
                "/api/auth/account",
                headers={"Authorization": f"Bearer {owner['access_token']}"},
            )
            assert deletion_response.status_code == 204
            assert not list((tmp_path / "uploads").glob("*"))
            expired_session = client.get(
                "/api/auth/me",
                headers={"Authorization": f"Bearer {owner['access_token']}"},
            )
            assert expired_session.status_code == 401
    finally:
        app.dependency_overrides.clear()
        Base.metadata.drop_all(engine)
        engine.dispose()
