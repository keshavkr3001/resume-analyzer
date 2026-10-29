from io import BytesIO

from docx import Document

from app.services.analysis import analyze_resume, match_job_description
from app.services.resume_parser import ResumeParseError, extract_resume_text


def test_analysis_reports_evidence_and_scoring():
    result = analyze_resume(
        "Jane Student\njane@example.com\nEducation\nBSc Computer Science\n"
        "Skills\nPython, SQL, pandas, Git\nProjects\nBuilt a data dashboard for 120 users."
    )
    assert 0 <= result["quality_score"] <= 100
    assert {skill["name"] for skill in result["extracted_skills"]} >= {"python", "sql", "pandas", "git"}
    assert all(skill["evidence"] for skill in result["extracted_skills"])
    assert "an ATS score" in result["career_recommendations"][0]["score_method"]


def test_job_match_is_transparent_and_bounded():
    result = match_job_description("Python SQL", "We need someone with Python and SQL, plus React.")
    assert result["match_score"] == 67
    assert result["matched_skills"] == ["python", "sql"]
    assert result["missing_skills"] == ["react"]


def test_parser_rejects_bad_pdf():
    try:
        extract_resume_text(b"not a pdf", ".pdf")
    except ResumeParseError as error:
        assert "valid PDF" in str(error)
    else:
        raise AssertionError("Invalid PDF should fail")


def test_docx_text_is_extracted():
    document = Document()
    document.add_paragraph("Education")
    document.add_paragraph("Computer science student with Python projects and experience.")
    buffer = BytesIO()
    document.save(buffer)
    extracted = extract_resume_text(buffer.getvalue(), ".docx")
    assert "Computer science student" in extracted
