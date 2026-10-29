from pathlib import Path

import pymupdf
from docx import Document


class ResumeParseError(ValueError):
    pass


def extract_resume_text(content: bytes, extension: str) -> str:
    try:
        if extension == ".pdf":
            if not content.startswith(b"%PDF-"):
                raise ResumeParseError("The uploaded file is not a valid PDF.")
            with pymupdf.open(stream=content, filetype="pdf") as document:
                text = "\n".join(page.get_text() for page in document)
        elif extension == ".docx":
            if not content.startswith(b"PK"):
                raise ResumeParseError("The uploaded file is not a valid DOCX document.")
            import io

            document = Document(io.BytesIO(content))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            for table in document.tables:
                text += "\n" + "\n".join(" ".join(cell.text for cell in row.cells) for row in table.rows)
        else:
            raise ResumeParseError("Only PDF and DOCX resumes are supported.")
    except ResumeParseError:
        raise
    except Exception as exc:
        raise ResumeParseError("The document could not be read. It may be damaged or password-protected.") from exc

    normalized = "\n".join(line.strip() for line in text.replace("\x00", "").splitlines() if line.strip())
    if len(normalized) < 40:
        raise ResumeParseError(
            "No usable text was found. This may be a scanned PDF; OCR is not configured."
        )
    return normalized[:100_000]
